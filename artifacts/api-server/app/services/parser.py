import io
import re


def _chunk_text(text: str, target: int = 800) -> list[str]:
    paragraphs = re.split(r"\n\n+", text)
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(para) > target:
            # Hard-split long paragraphs on word boundaries
            words = para.split()
            buf = ""
            for word in words:
                if len(buf) + len(word) + 1 > target and buf:
                    chunks.append(buf.strip())
                    buf = word
                else:
                    buf = (buf + " " + word).strip()
            if buf:
                chunks.append(buf)
            continue
        if len(current) + len(para) + 2 > target and current:
            chunks.append(current.strip())
            current = para
        else:
            current = (current + "\n\n" + para).strip() if current else para
    if current:
        chunks.append(current.strip())
    return [c for c in chunks if c]


async def extract_text(content: bytes, filename: str, content_type: str) -> tuple[str, str]:
    name_lower = filename.lower()

    if content_type.startswith("text/") or name_lower.endswith((".txt", ".md")):
        return content.decode("utf-8", errors="replace"), "text"

    if content_type == "application/pdf" or name_lower.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return text, "pdf"

    if name_lower.endswith(".docx") or content_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ):
        import docx
        doc = docx.Document(io.BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs)
        return text, "docx"

    if name_lower.endswith((".xlsx", ".xls")) or content_type in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        parts = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(",".join(str(c) if c is not None else "" for c in row))
            parts.append("\n".join(rows))
        return "\n\n".join(parts), "xlsx"

    # Fallback: treat as UTF-8 text
    return content.decode("utf-8", errors="replace"), "text"


def chunk_text(text: str, target: int = 800) -> list[str]:
    return _chunk_text(text, target)
